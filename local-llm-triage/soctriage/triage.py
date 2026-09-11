"""Glue - turn stage-3 outliers and Windows events into LLM items and write reports."""

import csv
from pathlib import Path

from .common import log, read_csv, write_csv, dump_json
from . import llm


def _outlier_to_text(row: dict) -> str:
    """Render one outlier row as a compact prompt line the model can judge."""
    art = row["artifact"]
    hc, th = row.get("host_count"), row.get("total_hosts")
    head = f"Artifact={art}; seen on host_count={hc} of total_hosts={th}; reason={row.get('reason','')}"
    detail = {
        "authorized_keys": ["user", "key_type", "comment", "options", "fingerprint"],
        "known_hosts": ["user", "key_type", "hostnames", "fingerprint"],
        "cron": ["run_as", "schedule", "command"],
        "services": ["unit", "user", "exec_start"],
        "users": ["username", "uid", "home", "shell", "flags"],
        "listening": ["proto", "port", "process", "local_addr"],
        "bodyfile": ["category", "path", "mode", "uid", "mtime"],
        "shell_history": ["user", "command", "flags"],
    }.get(art, [])
    parts = [f"{c}={row[c]}" for c in detail if row.get(c)]
    return head + ". " + "; ".join(parts) + f". hosts=[{row.get('hosts','')[:200]}]"


def triage_outliers(outliers_csv: Path, out_dir: Path, offline=False, model=llm.DEFAULT_MODEL,
                    url=llm.DEFAULT_URL, limit=None):
    """Read outliers.csv, triage each row with the local model, write triaged CSV + summary."""
    outliers_csv, out_dir = Path(outliers_csv), Path(out_dir)
    rows = list(read_csv(outliers_csv))
    if limit:
        rows = rows[:limit]
    items = [{"id": i, "text": _outlier_to_text(r), "_row": r} for i, r in enumerate(rows)]
    results = list(llm.triage_items(items, base_url=url, model=model, offline=offline))

    out_rows, counts = [], {}
    for res in results:
        r = dict(res["_row"])
        r.update({"verdict": res["verdict"], "confidence": res["confidence"],
                  "mitre": res["mitre"], "rationale": res["rationale"]})
        out_rows.append(r)
        counts[res["verdict"]] = counts.get(res["verdict"], 0) + 1
    order = {"malicious": 0, "suspicious": 1, "needs_review": 2, "benign": 3}
    out_rows.sort(key=lambda r: order.get(r.get("verdict"), 9))
    cols = ["verdict", "confidence", "mitre", "artifact", "reason", "host_count", "total_hosts",
            "rationale", "fingerprint", "exec_start", "command", "run_as", "path", "category",
            "username", "uid", "port", "process", "unit", "user", "hostnames",
            "key_type", "comment", "source_file", "hosts"]
    n = write_csv(out_dir / "triaged_outliers.csv", out_rows, cols)
    dump_json({"input": str(outliers_csv), "triaged": n, "verdicts": counts},
              out_dir / "triage_summary.json")
    log(f"triaged {n} outliers -> {out_dir/'triaged_outliers.csv'} | " +
        ", ".join(f"{k}={v}" for k, v in sorted(counts.items())))
    return counts


def triage_winlogs(xlsx_in: Path, xlsx_out: Path, offline=False, model=llm.DEFAULT_MODEL,
                   url=llm.DEFAULT_URL, notable_only=False, limit=None):
    """Review a Windows event-log Excel export with the local model; write reviewed workbook."""
    from . import winlogs
    xlsx_in, xlsx_out = Path(xlsx_in), Path(xlsx_out)
    headers, events, colmap = winlogs.load_events(xlsx_in, max_rows=limit)
    log(f"loaded {len(events)} events; columns detected: {sorted(colmap)}")

    idxs = winlogs.prefilter_notable(events) if notable_only else list(range(len(events)))
    items = [{"id": i, "text": winlogs.summarize_event(events[i])} for i in idxs]
    results = {r["id"]: r for r in llm.triage_items(items, base_url=url, model=model, offline=offline)}

    # Re-read the raw row values (aligned to sheet order) for the output workbook.
    from openpyxl import load_workbook
    wb = load_workbook(filename=str(xlsx_in), read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    next(it, None)  # header
    raw_rows = []
    for i, r in enumerate(it):
        if limit and i >= limit:
            break
        raw_rows.append(list(r))
    wb.close()

    default = {"verdict": "benign", "confidence": "low", "rationale": "not selected for review", "mitre": ""}
    verdicts = [results.get(i, default) for i in range(len(raw_rows))]
    stats = winlogs.write_reviewed_xlsx(headers, raw_rows, verdicts, xlsx_out)
    log(f"reviewed {stats['total']} events, flagged {stats['flagged']} -> {stats['out']}")
    return stats
